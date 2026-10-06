"""検索候補は人の採用まで分析に混ぜない。"""
import re
from datetime import datetime, timezone
from collections import Counter
from openpyxl import load_workbook
from core.search_state import identity
from core.deduplicate import group_by_asin
from keepa.parser import basic_product_fields,product_codes
from matchers.product_identity import evaluate,search_terms

def read_reviews(path):
    if not path:return []
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        if '確認入力' not in wb.sheetnames:raise ValueError('確認入力シートがありません')
        it=wb['確認入力'].iter_rows(values_only=True);headers=next(it,())
        if not {'source_id','input_hash','ASIN','採用'}.issubset(headers):raise ValueError('確認入力の列が不足しています')
        out=[]
        for row in it:
            r=dict(zip(headers,row));decision=str(r.get('採用') or '').strip()
            if decision not in ('','採用','除外','保留'):raise ValueError('採用欄は採用／除外／保留を指定してください')
            if decision:
                if not re.fullmatch(r'[A-Z0-9]{10}',str(r.get('ASIN') or '')):raise ValueError('確認ASINが不正です')
                r['採用']=decision;out.append(r)
        return out
    finally:wb.close()

def lookup_by_name(products,matching,client,state,*,limit=20,query_limit=2,token_budget=400,
                   include_no_candidate=False,force=False,reviews=()):
    if any(type(v) is not int or v<0 for v in [limit,token_budget]) or type(query_limit) is not int or query_limit not in (1,2):raise ValueError('商品名検索設定が不正です')
    out={'candidates':[],'statuses':[],'errors':[],'approved':[],'summary':{}}
    targets=[r for r in matching['unresolved'] if r['match_status']=='INSUFFICIENT_DATA' or (include_no_candidate and r['match_status']=='NO_CANDIDATE')]
    seen={}
    for p in products:
        s,h=identity(p);seen.setdefault(s,set()).add(h)
    def ids(p):
        s,h=identity(p);return (s+'-'+h[:12] if len(seen[s])>1 else s),h
    valid={ids(p) for p in products};pending={}
    for r in reviews:
        s,h,a=r.get('source_id'),r.get('input_hash'),r.get('ASIN')
        if (s,h) not in valid:raise ValueError('確認Excelの入力商品が変更されています')
        if (s,h,a) in pending and pending[s,h,a]!=r['採用']:raise ValueError('確認Excelに矛盾する採用があります')
        pending[s,h,a]=r['採用']
    used=fresh=queries=0
    for base in targets:
        p=base['product'];sid,fp=ids(p);terms=search_terms(p,query_limit);candidates={}
        status='候補なし';reason='今回の検索で候補なし。Amazon不存在は断定しません';new=False
        for term in terms:
            key,cached=state.query(term,force=force)
            if cached is None:
                if (limit and fresh>=limit and not new) or (token_budget and used+10>token_budget):
                    status='未調査';reason='件数または検索予算の上限';break
                if not new:fresh+=1;new=True
                try:
                    before=client.total_tokens_consumed
                    data=client.keepa_get_with_wait('/search',{'domain':5,'type':'product','term':term,'history':0,'stats':90,'update':48},10,'商品名検索')
                    cached=data.get('products')
                    if not isinstance(cached,list) or any(not isinstance(v,dict) or not re.fullmatch(r'[A-Z0-9]{10}',str(v.get('asin',''))) for v in cached):raise ValueError('検索応答不正')
                except Exception:
                    used=token_budget or used+10;status='通信エラー';reason='取得失敗。候補なしとは判定しません'
                    out['errors'].append({'source_id':sid,'処理':'商品名検索','理由':reason});break
                used+=max(10,client.total_tokens_consumed-before);queries+=1
                state.put(key,cached)
            for c in cached:
                a=c['asin']
                if a not in candidates:
                    candidates[a]={'source_id':sid,'input_hash':fp,'source_index':base['source_index'],
                        '入力商品名':p.get('product_name'),'入力ブランド':p.get('brand'),'入力容量':p.get('capacity'),
                        '原入力JAN':p.get('jan'),'ASIN':a,'Amazon商品名':c.get('title'),'候補側コード':sorted(product_codes(c)),
                        '検索語':term,'確認日時':datetime.now(timezone.utc).isoformat(),'Amazon商品ページ':'https://www.amazon.co.jp/dp/'+a,**evaluate(p,c),'_amazon':c}
            if any(r['評価']=='有力候補・未承認' for r in candidates.values()):break
        if not terms:status='入力不足';reason='商品名とブランドまたは型番が必要'
        elif status not in ('未調査','通信エラー') and candidates:status='確認待ち';reason='採用欄で人の確認が必要'
        prior=state.data['reviews'].get(sid,{})
        decisions=dict(prior.get('decisions',{})) if prior.get('hash')==fp else {}
        for (s,h,a),d in list(pending.items()):
            if s==sid and h==fp:
                if a not in candidates:raise ValueError('確認ASINがこの商品の検索履歴にありません。再検索結果を確認してください')
                decisions[a]=d;del pending[s,h,a]
        adopted=[a for a,d in decisions.items() if d=='採用' and a in candidates]
        if len(adopted)>1:raise ValueError('同じ商品への複数ASIN採用は保留してください')
        state.data['reviews'][sid]={'hash':fp,'decisions':decisions}
        for a,r in candidates.items():
            amazon=r.pop('_amazon');r['採用']=decisions.get(a,'保留');out['candidates'].append(r)
            if a in adopted:out['approved'].append({**base,'asin':a,'amazon':basic_product_fields(amazon),'match_method':'NAME_REVIEWED','match_status':'HUMAN_APPROVED'})
        if adopted:status='採用済み';reason='確認入力で採用'
        out['statuses'].append({'source_id':sid,'input_hash':fp,'商品名':p.get('product_name'),'状態':status,'理由':reason,'候補ASIN数':len(candidates)})
    if pending:raise ValueError('確認Excelに今回の対象外の行があります')
    state.save();counts=Counter(r['状態'] for r in out['statuses'])
    out['summary']={'name_target_products':len(targets),'name_new_products':fresh,'name_api_queries':queries,
        'name_tokens_estimated_or_actual':used,'name_candidate_rows':len(out['candidates']),'name_approved_products':len(out['approved']),**{'name_'+k:v for k,v in counts.items()}}
    return out

def merge_approved(matching,names):
    matching['matched'].extend(names['approved']);matching['asin_products']=group_by_asin(matching['matched'])
    approved={r['source_index'] for r in names['approved']};matching['unresolved']=[r for r in matching['unresolved'] if r['source_index'] not in approved]
    matching['summary'].update(names['summary']);matching['summary'].update(matched_source_products=len({r['source_index'] for r in matching['matched']}),matched_rows=len(matching['matched']),unique_asins=len(matching['asin_products']),unresolved_products=len(matching['unresolved']))


def refresh_approved(names, client):
    """採用済みの価格等は検索キャッシュを流用せず取得する。"""
    rows=names['approved'];asins=list(dict.fromkeys(r['asin'] for r in rows));received={}
    for offset in range(0,len(asins),20):
        batch=asins[offset:offset+20]
        try:
            data=client.keepa_get_with_wait('/product',{'domain':5,'asin':','.join(batch),'stats':90,'history':0,'update':1},len(batch),'採用済み商品取得')
            products=data.get('products')
            if not isinstance(products,list) or any(not isinstance(p,dict) for p in products):raise ValueError('応答不正')
            for p in products:
                if p.get('asin') in batch:received[p['asin']]=p
        except Exception:
            names['errors'].append({'処理':'採用済み商品取得','ASIN':batch,'理由':'取得失敗。今回の分析対象外'})
    names['approved']=[]
    for r in rows:
        if r['asin'] in received:
            r['amazon']=basic_product_fields(received[r['asin']]);names['approved'].append(r)
        else:
            names['errors'].append({'処理':'採用済み商品取得','ASIN':r['asin'],'理由':'応答に商品なし。今回の分析対象外'})
    names['summary']['name_approved_products']=len(names['approved'])
