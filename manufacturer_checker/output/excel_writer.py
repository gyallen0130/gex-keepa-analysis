"""既存v4.4のopenpyxl出力をメーカー共通形式へ移植。"""
import json
import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def flatten(row):
    if 'product' in row:
        return {**row['product'], **{k:v for k,v in row.items() if k != 'product'}}
    return row

def write_excel(path, products, matching, analysis, *, collector_errors=(), tokens_consumed=None, input_warnings=()):
    rows = analysis['rows']
    if len({r['asin'] for r in rows}) != len(rows):
        raise ValueError('ASIN重複が残っています')
    flat_rows = [{**(row['links'][0]['product'] if row['links'] else {}),
                  **{k:v for k,v in row.items() if k not in ('amazon','links')}} for row in rows]
    unresolved = [flatten(r) for r in matching['unresolved']]
    links = []
    for row in rows:
        for link in row['links']:
            links.append({'asin':row['asin'], **flatten(link)})
    summary = [{'項目':k,'値':v} for k,v in matching['summary'].items()]
    summary += [{'項目':k,'値':v} for k,v in analysis['settings'].items()]
    summary += [{'項目':'候補ASIN数','値':sum(r['候補']=='○' for r in rows)}, {'項目':'Keepa消費トークン','値':tokens_consumed},
        {'項目':'JAN照合','値':'識別コード一致のみ。単品/セット・容量等は未検証'},
        {'項目':'判定閾値','値':'min_other_shareは0〜100の百分率。0.1は0.1%です'},
        {'項目':'未照合','値':'Amazonに存在しないとは断定しません'}]
    sheets = {'サマリー':summary,'候補_全て':[r for r in flat_rows if r['候補']=='○'],
        'Amazon照合済み':flat_rows,'メーカー商品マスタ':products,'未照合':[r for r in unresolved if r['match_status']!='NOT_SCANNED'],
        '今回未調査':[r for r in unresolved if r['match_status']=='NOT_SCANNED'],
        'BB未確認':[r for r in flat_rows if r['判定'].startswith(('D：','F：'))],
        'ASIN商品対応':links,'商品取得エラー':list(collector_errors), '入力確認':list(input_warnings),
        'Keepa取得エラー':matching['errors']+analysis['errors']}
    wb=Workbook();wb.remove(wb.active)
    for name, records in sheets.items():
        ws=wb.create_sheet(name)
        headers=list(dict.fromkeys(k for record in records for k in record)) or ['情報']
        ws.append(headers)
        for record in records:
            values=[]
            for key in headers:
                value=record.get(key)
                if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
                if isinstance(value,float) and not math.isfinite(value):value=None
                values.append(value)
            ws.append(values)
            # 外部商品名が数式として実行されないよう文字列を明示。
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78');cell.alignment=Alignment(wrap_text=True)
        for idx,header in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(idx)].width=min(45,max(12,len(header)+2))
            if header in ('jan','manufacturer_code','asin','ASIN'):
                for cell in list(ws.columns)[idx-1][1:]:cell.number_format='@'
        for cells in ws.iter_rows(min_row=2):
            values={header:cell.value for header,cell in zip(headers,cells)}
            shade='E2F0D9' if values.get('候補')=='○' else 'FFF2CC' if str(values.get('判定','')).startswith(('D：','F：')) else None
            for header,cell in zip(headers,cells):
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if shade:cell.fill=PatternFill('solid',fgColor=shade)
                if header in ('official_url','source_url','Amazon商品ページ') and isinstance(cell.value,str) and cell.value.startswith(('https://','http://')):
                    cell.hyperlink=cell.value;cell.font=Font(color='0563C1',underline='single')
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
    return path
