"""Excel/CSV商品マスタを共通形式へ変換。元の行と読み込み警告を保持。"""
import csv
import math
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook

FIELDS = ('manufacturer','brand','product_name','jan','model_number','manufacturer_code',
          'category','capacity','size','color','pack_quantity','official_url',
          'search_keywords','unit_capacity','sales_units','variant')
ALIASES = {
 'manufacturer':['メーカー','メーカー名','製造元'], 'brand':['ブランド','ブランド名'],
 'product_name':['商品名','品名','製品名','公式商品名'],
 'jan':['JAN','JANコード','JAN/EAN','EAN','EANコード','バーコード'],
 'model_number':['型番','メーカー型番'], 'manufacturer_code':['商品コード','品番','メーカー商品コード'],
 'search_keywords':['検索キーワード'], 'unit_capacity':['単品容量'],
 'sales_units':['販売単位数'], 'variant':['味・種類'],
 'category':['カテゴリ','カテゴリー'], 'capacity':['容量','内容量'], 'size':['サイズ'],
 'color':['色','カラー'], 'pack_quantity':['入数','入り数','数量'],
 'official_url':['公式URL','公式商品URL','公式商品ページ'],
}

def _header(value):
    return re.sub(r'[\s_]+','',unicodedata.normalize('NFKC',str(value or ''))).casefold()

def _text(value):
    return '' if value is None else str(value).strip()

def _jan(value):
    if value is None or value == '':
        return '', ''
    if isinstance(value, bool):
        return '', 'JANの真偽値は無効'
    if isinstance(value, int):
        text = str(value)
        warning = 'JANが数値セル。先頭ゼロが失われていないか確認してください'
    elif isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return '', 'JANが小数/非有限数で変換不可'
        text, warning = str(int(value)), 'JANが数値セル。先頭ゼロが失われていないか確認してください'
    else:
        text, warning = unicodedata.normalize('NFKC',str(value)).strip(), ''
    text = re.sub(r'[\s-]','',text)
    if not re.fullmatch(r'[0-9]{8,14}',text):
        return '', 'JAN形式が不正（指数表記や桁不足は自動補完しません）'
    return text, warning

class ExcelCollector:
    def __init__(self, path, manufacturer='', *, sheet_name=None, header_row=1, column_map=None, encoding=None):
        self.path = Path(path)
        self.manufacturer = manufacturer
        self.sheet_name = sheet_name
        if not isinstance(header_row,int) or header_row < 1:
            raise ValueError('header_rowは1以上')
        self.header_row = header_row
        self.column_map = column_map or {}
        if any(key not in FIELDS for key in self.column_map):
            raise ValueError('column_mapのキーが共通項目にありません')
        self.encoding = encoding
        self.errors = []
        self.warnings = []

    def collect(self):
        self.errors, self.warnings = [], []
        suffix = self.path.suffix.lower()
        workbook = None
        if suffix == '.xlsx':
            workbook = load_workbook(self.path, read_only=True, data_only=False)
            if self.sheet_name and self.sheet_name not in workbook.sheetnames:
                workbook.close();raise ValueError(f'指定シートがありません: {self.sheet_name}')
            sheet = workbook[self.sheet_name] if self.sheet_name else workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            source_sheet = sheet.title
        elif suffix in ('.csv','.tsv'):
            raw = self.path.read_bytes()
            if self.encoding:
                text = raw.decode(self.encoding)
            else:
                try: text = raw.decode('utf-8-sig')
                except UnicodeDecodeError: text = raw.decode('cp932')
            import io
            rows = iter(csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','))
            source_sheet = ''
        else:
            raise ValueError('対応形式は.xlsx/.csv/.tsvです（.xlsは.xlsxへ保存し直してください）')
        try:
            header = None
            for _ in range(self.header_row):
                header = next(rows, None)
            if header is None:
                raise ValueError('見出し行がありません')
            normalized = [_header(value) for value in header]
            mapping = {}
            for field in FIELDS:
                names = [self.column_map[field]] if field in self.column_map else [field]+ALIASES.get(field, [])
                options = {_header(name) for name in names}
                indexes = [i for i,value in enumerate(normalized) if value in options]
                if len(indexes) > 1:
                    raise ValueError(f'{field}の候補列が複数あります。column_mapで指定してください')
                if field in self.column_map and not indexes:
                    raise ValueError(f'指定列がありません: {self.column_map[field]}')
                if indexes: mapping[field] = indexes[0]
            if 'product_name' not in mapping and 'jan' not in mapping:
                raise ValueError('商品名またはJAN列が必要です')
            products = []
            for row_number, row in enumerate(rows, self.header_row+1):
                if all(value is None or str(value).strip()=='' for value in row):
                    continue
                product = {field:'' for field in FIELDS}
                for field,index in mapping.items():
                    value = row[index] if index < len(row) else None
                    if isinstance(value,str) and value.startswith('='):
                        self.warnings.append({'行':row_number,'項目':field,'理由':'数式は評価しません。値に置換してください'})
                        value = None
                    if field == 'jan':
                        product[field], warning = _jan(value)
                        if warning:self.warnings.append({'行':row_number,'項目':field,'理由':warning})
                    else:product[field] = _text(value)
                if not product['manufacturer']:product['manufacturer']=str(self.manufacturer).strip()
                product.update(source_type='file', source_url='', source_file=self.path.name,
                    source_sheet=source_sheet, source_row=row_number)
                products.append(product)
            return products
        finally:
            if workbook is not None:workbook.close()
