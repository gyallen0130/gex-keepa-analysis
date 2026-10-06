import tempfile,unittest
from pathlib import Path
from openpyxl import load_workbook
from core.search_state import SearchState
from matchers.jan_matcher import lookup_by_jan
from matchers.product_identity import evaluate,quantities
from matchers.name_matcher import lookup_by_name,read_reviews
from ui.colab_runner import run_research
P={'manufacturer':'日仏商事','manufacturer_code':'10802000','brand':'サバトン','product_name':'マロンペースト','capacity':'1kg','sales_units':'1','jan':''}
C={'asin':'B012345678','title':'サバトン マロンペースト 1000g 1個','brand':'サバトン','eanList':['3101740000446'],'stats':{'current':[-1,2000]}}
class Client:
    def __init__(self,responses):self.responses=iter(responses);self.calls=[];self.total_tokens_consumed=0
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def keepa_get_with_wait(self,endpoint,params,*args):
        self.calls.append((endpoint,params));r=next(self.responses)
        if isinstance(r,Exception):raise r
        self.total_tokens_consumed+=10;return r
class Tests(unittest.TestCase):
    def test_capacity_and_pack(self):
        self.assertEqual(quantities('1kg'),quantities('1000g'))
        self.assertEqual(evaluate(P,C)['評価'],'有力候補・未承認')
        self.assertEqual(evaluate(P,{**C,'title':'サバトン マロンペースト 500g×2個'})['評価'],'別規格・除外')
        self.assertEqual(evaluate(P,{**C,'title':'サバトン マロンペースト 1kg'})['比較結果']['販売単位'],'不明')
        self.assertEqual(evaluate(P,{**C,'title':'サバトン マロンペースト 1L 1個'})['比較結果']['容量'],'不一致')
    def test_review_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            client=Client([{'products':[C]}]);state=Path(d)/'state.json'
            path,_,report=run_research([P],'secret','日仏商事',enable_name_search=True,bb_limit=0,state_path=state,output_dir=d,client_factory=lambda _:client)
            self.assertEqual(report['unique_asins'],0);self.assertEqual(report['name_candidate_rows'],1)
            wb=load_workbook(path);ws=wb['確認入力'];headers=[c.value for c in ws[1]]
            ws.cell(2,headers.index('採用')+1,'採用');wb.save(path);wb.close()
            self.assertEqual(read_reviews(path)[0]['採用'],'採用')
            client2=Client([{'products':[C]}])
            path2,_,r=run_research([P],'secret','日仏商事',enable_name_search=True,review_path=path,bb_limit=0,state_path=state,output_dir=d,client_factory=lambda _:client2)
            self.assertEqual(r['unique_asins'],1);self.assertTrue(all(endpoint=='/product' for endpoint,_ in client2.calls))
            wb=load_workbook(path2);self.assertIn('NAME_REVIEWED',str(list(wb['ASIN商品対応'].values)));wb.close()
            self.assertNotIn('secret',state.read_text())
    def test_budget_resume(self):
        with tempfile.TemporaryDirectory() as d:
            state=SearchState(Path(d)/'state.json');p={**P,'sales_units':''};m=lookup_by_jan([p],Client([]))
            a=lookup_by_name([p],m,Client([{'products':[C]}]),state,token_budget=10)
            self.assertEqual(a['statuses'][0]['状態'],'未調査')
            c=Client([{'products':[]}]);b=lookup_by_name([p],m,c,SearchState(state.path))
            self.assertEqual(len(c.calls),1);self.assertEqual(b['statuses'][0]['状態'],'確認待ち')
    def test_no_error_or_limit_bypass(self):
        products=[P,{**P,'jan':'12345678'},{**P,'jan':'87654321'}]
        m={'unresolved':[{'source_index':i,'product':p,'match_status':s} for i,(p,s) in enumerate(zip(products,['INSUFFICIENT_DATA','LOOKUP_ERROR','NOT_SCANNED']))]}
        with tempfile.TemporaryDirectory() as d:
            r=lookup_by_name(products,m,Client([{'products':[C]}]),SearchState(Path(d)/'s.json'),include_no_candidate=True)
            self.assertEqual(len(r['statuses']),1)
    def test_error_and_sync_failure(self):
        with tempfile.TemporaryDirectory() as d:
            s=SearchState(Path(d)/'s.json');m=lookup_by_jan([P],Client([]))
            r=lookup_by_name([P],m,Client([RuntimeError('secret')]),s)
            self.assertEqual(r['statuses'][0]['状態'],'通信エラー');self.assertEqual(s.data['queries'],{})
            def bad(_):raise OSError('sync failed')
            with self.assertRaises(OSError):lookup_by_name([P],m,Client([{'products':[C]}]),SearchState(s.path,sync=bad))
    def test_changed_input_or_foreign_candidate(self):
        with tempfile.TemporaryDirectory() as d:
            s=SearchState(Path(d)/'s.json');m=lookup_by_jan([P],Client([]));a=lookup_by_name([P],m,Client([{'products':[C]}]),s)
            r={**a['candidates'][0],'採用':'採用'};p={**P,'capacity':'500g'}
            with self.assertRaises(ValueError):lookup_by_name([p],lookup_by_jan([p],Client([])),Client([]),s,reviews=[r])
            r['ASIN']='B999999999'
            with self.assertRaises(ValueError):lookup_by_name([P],m,Client([]),s,reviews=[r])
    def test_disabled(self):
        with tempfile.TemporaryDirectory() as d:
            c=Client([]);_,_,r=run_research([P],'k','日仏商事',output_dir=d,client_factory=lambda _:c)
            self.assertEqual(c.calls,[]);self.assertEqual(r['unique_asins'],0)
