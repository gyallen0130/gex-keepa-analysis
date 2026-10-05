import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from analysis.buybox import calculate_bb_share
from analysis.candidate_judge import judge_row
from analysis.service import analyze_buybox
from matchers.jan_matcher import lookup_by_jan
from output.excel_writer import write_excel
from tests.test_matcher import Fake
class Tests(unittest.TestCase):
 def test_history_boundaries(self):
  now=200000;start=now-90*1440
  r=calculate_bb_share([start-10,'AMZ',start+45*1440,'OTHER'],90,{'AMZ'},now=now)
  self.assertEqual(r['amazon_pct'],50);self.assertEqual(r['other_pct'],50);self.assertEqual(r['coverage_pct'],100)
  r=calculate_bb_share([now-1440,'OTHER'],90,{'AMZ'},now=now)
  self.assertEqual(r['other_pct'],100);self.assertLess(r['coverage_pct'],2)
  self.assertIsNone(calculate_bb_share([],90,{'AMZ'},now=now)['other_pct'])
 def test_judge_coverage_first(self):
  row={'Amazon現在':'あり','現在カート':'Amazon','90日_他セラーカート率%':100,'90日_履歴カバー率%':1}
  self.assertTrue(judge_row(row)[0].startswith('F：'))
  row['90日_履歴カバー率%']=100;row['90日_他セラーカート率%']=0.2
  self.assertTrue(judge_row(row)[0].startswith('C：'))
  self.assertTrue(judge_row(row,min_other_share=10)[0].startswith('E：'))
  self.assertEqual(judge_row({})[1],'')
 def test_missing_stats_is_unknown(self):
  grouped=[{'asin':'A','amazon':{'Amazon現在':'あり'},'links':[]}]
  r=analyze_buybox(grouped,Fake([{'products':[{'asin':'A'}]}]),now=200000)
  self.assertEqual(r['rows'][0]['現在カート'],'セラー不明');self.assertEqual(r['rows'][0]['候補'],'')
 def test_end_to_end_workbook(self):
  now=200000
  ps=[{'manufacturer':'別メーカー','product_name':'=1+1','jan':'12345678'},{'manufacturer':'GEX','jan':'12345678'},{'jan':''}]
  matching=lookup_by_jan(ps,Fake([{'products':[{'asin':'A','eanList':['12345678'],'stats':{'current':[100000]}}]}]))
  analysis=analyze_buybox(matching['asin_products'],Fake([{'products':[{'asin':'A','stats':{'buyBoxSellerId':'OTHER','buyBoxIsAmazon':False},'buyBoxSellerIdHistory':[now-90*1440,'OTHER']}]}]),now=now)
  self.assertTrue(analysis['rows'][0]['判定'].startswith('B：'))
  with tempfile.TemporaryDirectory() as d:
   p=write_excel(Path(d)/'test.xlsx',ps,matching,analysis)
   wb=load_workbook(p)
   self.assertEqual(wb['候補_全て'].max_row,2);self.assertEqual(wb['ASIN商品対応'].max_row,3)
   self.assertEqual(wb['メーカー商品マスタ']['B2'].data_type,'s')
   self.assertEqual(wb['Amazon照合済み'].freeze_panes,'A2')
   self.assertEqual(wb['メーカー商品マスタ']['C2'].value,'12345678')
 def test_unknown_amazon_not_candidate(self):
  from keepa.parser import basic_product_fields
  row=basic_product_fields({"asin":"A"})
  self.assertEqual(row["Amazon現在"],"未確認")
  self.assertEqual(judge_row(row)[1],"")
 def test_skip_and_error(self):
  grouped=[{'asin':'A','amazon':{'Amazon現在':'あり'},'links':[]}]
  self.assertTrue(analyze_buybox(grouped,Fake([]),limit=0)['rows'][0]['判定'].startswith('D：'))
  result=analyze_buybox(grouped,Fake([RuntimeError('fail')]))
  self.assertEqual(result['rows'][0]['BB取得状態'],'ERROR');self.assertEqual(result['rows'][0]['候補'],'')
