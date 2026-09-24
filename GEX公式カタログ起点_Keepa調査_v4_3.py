# GEX公式カタログ起点 Keepa調査 v4.3
# Google Colab用。Notebook版では先頭セルで依存パッケージを導入します。

import re
import time
import math
import getpass
import threading
from collections import deque, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from urllib.parse import urljoin, urlparse, parse_qs, urlencode, urlunparse

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup
from tqdm.auto import tqdm
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    from google.colab import files
    IN_COLAB = True
except ImportError:
    files = None
    IN_COLAB = False

try:
    from IPython.display import display
except ImportError:
    def display(value):
        print(value)


# ================================================================
# 1. 設定
# ================================================================

OFFICIAL_ROOTS = {
    "アクアリウム": "https://product.gex-fp.co.jp/fish/",
    "犬猫": "https://product.gex-fp.co.jp/ca/",
    "小動物": "https://product.gex-fp.co.jp/animal/",
    "爬虫類（エキゾテラ）": "https://product.gex-fp.co.jp/exoterra/",
}

DOMAIN = 5
AMAZON_JP_SELLER_ID = "AN1VRQENFRJN5"

# 公式サイトへの同時アクセス数。先方サイトへの負荷を抑えるため4以下を推奨。
OFFICIAL_WORKERS = 4
REQUEST_TIMEOUT = 45
REQUEST_RETRIES = 4

# Keepaは小分けに取得し、足りない場合は補充を待ってから進める。
KEEPA_CODE_BATCH = 20
KEEPA_ASIN_BATCH = 20
TOKEN_WAIT_BUFFER = 5
DEFAULT_REFILL_RATE = 20

# 他セラーの90日カート率がこの値以上なら候補。
MIN_OTHER_BB_SHARE_90 = 0.1
MIN_BB_COVERAGE_PERCENT = 50.0

OUTPUT_FILE = "GEX公式全商品_Keepa調査_v4_3.xlsx"
OFFICIAL_ONLY_FILE = "GEX公式商品マスタ_v4.xlsx"

USER_AGENT = (
    "Mozilla/5.0 (compatible; GEXOfficialCatalogResearch/4.3; "
    "+https://product.gex-fp.co.jp/)"
)

_HTTP_LOCAL = threading.local()


# ================================================================
# 2. 共通関数
# ================================================================

def canonicalize_url(url):
    """フラグメントを除き、クエリ順を統一する。"""
    p = urlparse(url)
    query = parse_qs(p.query, keep_blank_values=True)
    flat = []
    for key in sorted(query):
        for value in sorted(query[key]):
            flat.append((key, value))
    return urlunparse((p.scheme, p.netloc, p.path, "", urlencode(flat), ""))


def request_html(url, retry=REQUEST_RETRIES):
    if not hasattr(_HTTP_LOCAL, "session"):
        _HTTP_LOCAL.session = requests.Session()
        _HTTP_LOCAL.session.headers.update({
            "User-Agent": USER_AGENT,
            "Accept-Language": "ja",
            "Accept-Encoding": "gzip, deflate",
        })
    last_error = None
    for attempt in range(retry):
        try:
            r = _HTTP_LOCAL.session.get(
                url,
                timeout=REQUEST_TIMEOUT,
            )
            r.raise_for_status()
            r.encoding = r.apparent_encoding or r.encoding
            return r.text
        except Exception as e:
            last_error = e
            if attempt < retry - 1:
                time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"公式ページ取得失敗: {url} / {last_error}")


def query_value(url, key):
    values = parse_qs(urlparse(url).query).get(key, [])
    return values[0] if values else None


def is_same_section(url, root_url):
    u = urlparse(url)
    r = urlparse(root_url)
    return (
        u.scheme in ("http", "https")
        and u.netloc == r.netloc
        and u.path.startswith(r.path)
    )


def is_list_url(url, root_url):
    return is_same_section(url, root_url) and query_value(url, "m") == "ProductList"


def is_detail_url(url, root_url):
    return (
        is_same_section(url, root_url)
        and query_value(url, "m") == "ProductListDetail"
        and query_value(url, "id") is not None
    )


def valid_jan(value):
    value = re.sub(r"\D", "", str(value or ""))
    return value if 8 <= len(value) <= 14 else None


def normalize_code(value):
    value = re.sub(r"\D", "", str(value or ""))
    return value.lstrip("0") or value


def safe_number(value):
    try:
        if value is None or (isinstance(value, float) and math.isnan(value)):
            return None
        return float(value)
    except Exception:
        return None


def safe_array_get(array, index):
    if not isinstance(array, list) or index >= len(array):
        return None
    value = array[index]
    if value in (-1, -2, None):
        return None
    return value


def yen_from_keepa(value):
    value = safe_number(value)
    if value is None or value < 0:
        return None
    return round(value / 100, 2)


# ================================================================
# 3. 公式サイト：商品一覧URLと商品詳細URLの収集
# ================================================================

def discover_detail_urls(section_name, root_url):
    queue = deque([canonicalize_url(root_url)])
    visited_lists = set()
    detail_urls = set()
    errors = []

    while queue:
        # 現在見つかっている一覧ページをまとめて並列取得する。
        batch = []
        while queue:
            url = queue.popleft()
            if url not in visited_lists:
                visited_lists.add(url)
                batch.append(url)

        with ThreadPoolExecutor(max_workers=OFFICIAL_WORKERS) as executor:
            futures = {executor.submit(request_html, url): url for url in batch}
            for future in as_completed(futures):
                url = futures[future]
                try:
                    html = future.result()
                except Exception as e:
                    errors.append({"区分": section_name, "URL": url, "エラー": str(e)})
                    continue

                soup = BeautifulSoup(html, "html.parser")
                for a in soup.select("a[href]"):
                    absolute = canonicalize_url(urljoin(url, a.get("href")))
                    if is_detail_url(absolute, root_url):
                        detail_urls.add(absolute)
                    elif is_list_url(absolute, root_url) and absolute not in visited_lists:
                        queue.append(absolute)

    return sorted(detail_urls), errors, len(visited_lists)


def parse_official_product(section_name, url):
    html = request_html(url)
    soup = BeautifulSoup(html, "html.parser")

    h1 = (
        soup.select_one("h1.c-detail__ttl")
        or soup.select_one("h1.c-detail__ttl__sp")
        or soup.select_one("h1.page-header__title")
    )
    product_name = h1.get_text(" ", strip=True) if h1 else ""
    if not product_name and soup.title:
        product_name = re.sub(
            r"\s*[|｜]\s*ジェックス株式会社\s*$",
            "",
            soup.title.get_text(" ", strip=True),
        ).strip()

    crumbs = []
    crumb = soup.select_one("p.pankuzu")
    if crumb:
        crumbs = [x.strip() for x in crumb.get_text("|", strip=True).split("|") if x.strip()]
        crumbs = [x for x in crumbs if x != "＞"]

    specs = {}
    for tr in soup.select("tr"):
        th = tr.find("th")
        td = tr.find("td")
        if th and td:
            specs[th.get_text(" ", strip=True)] = td.get_text(" ", strip=True)

    jan_text = specs.get("JANコード", "")
    jans = sorted(set(re.findall(r"\d{8,14}", jan_text)))
    if not jans:
        page_text = soup.get_text("\n", strip=True)
        match = re.search(r"JANコード\s*([0-9]{8,14})", page_text)
        if match:
            jans = [match.group(1)]

    product_id = query_value(url, "id")
    cid = query_value(url, "cid")
    category_path = " ＞ ".join(crumbs[2:-1]) if len(crumbs) >= 4 else ""

    base = {
        "公式区分": section_name,
        "公式商品名": product_name,
        "公式カテゴリ": category_path,
        "公式商品ID": product_id,
        "公式カテゴリID": cid,
        "商品コード": specs.get("コード"),
        "標準小売価格": specs.get("標準小売価格"),
        "原産国": specs.get("原産国"),
        "公式商品ページ": url,
    }

    if jans:
        return [{**base, "JANコード": jan} for jan in jans]
    return [{**base, "JANコード": None}]


print("=" * 72)
print("STEP 1：GEX公式商品マスタの準備")
print("=" * 72)
print("2回目以降は、以前作成した次のどちらかを再利用できます。")
print("・GEX公式商品マスタ_v4.xlsx")
print("・GEX公式全商品_Keepa調査_v4.xlsx")
reuse_master = input("既存ファイルを再利用しますか？（推奨 YES／初回のみ NO）：").strip().upper()

MASTER_REUSED = reuse_master == "YES"
master_source_file = None

if MASTER_REUSED:
    if IN_COLAB:
        print("\n既存のマスタまたは前回結果Excelをアップロードしてください。")
        uploaded = files.upload()
        excel_files = [
            name for name in uploaded.keys()
            if name.lower().endswith((".xlsx", ".xlsm"))
        ]
        if not excel_files:
            raise RuntimeError("Excelファイルがアップロードされていません。")
        master_source_file = excel_files[0]
    else:
        master_source_file = input("既存Excelのパス：").strip()

    xls = pd.ExcelFile(master_source_file)
    if "公式商品マスタ" not in xls.sheet_names:
        raise RuntimeError("アップロードしたExcelに『公式商品マスタ』シートがありません。")

    df_official = pd.read_excel(
        master_source_file,
        sheet_name="公式商品マスタ",
        dtype={"JANコード": str},
    )
    required_master_columns = {"公式区分", "公式商品ページ", "JANコード"}
    missing_columns = required_master_columns - set(df_official.columns)
    if missing_columns:
        raise RuntimeError(
            "公式商品マスタの必要列が不足しています："
            + ", ".join(sorted(missing_columns))
        )

    df_official_errors = pd.DataFrame(columns=["区分", "URL", "エラー"])
    print(f"\n既存マスタを読み込みました：{master_source_file}")

else:
    print("\nGEX公式サイトから商品マスタを新規作成します。")
    print("対象区分：", " / ".join(OFFICIAL_ROOTS.keys()))
    print("※ エキゾテラはGEX公式の爬虫類用品区分として含めます。\n")

    all_detail_tasks = []
    official_errors = []
    list_page_counts = {}

    for section_name, root_url in OFFICIAL_ROOTS.items():
        print(f"[{section_name}] 一覧ページを探索中...")
        urls, errs, list_count = discover_detail_urls(section_name, root_url)
        all_detail_tasks.extend((section_name, url) for url in urls)
        official_errors.extend(errs)
        list_page_counts[section_name] = list_count
        print(f"  一覧ページ {list_count} / 商品詳細URL {len(urls)}")

    all_detail_tasks = list(dict.fromkeys(all_detail_tasks))
    official_rows = []

    print(f"\n商品詳細ページを取得します：{len(all_detail_tasks)}ページ")
    with ThreadPoolExecutor(max_workers=OFFICIAL_WORKERS) as executor:
        futures = {
            executor.submit(parse_official_product, section, url): (section, url)
            for section, url in all_detail_tasks
        }
        for future in tqdm(as_completed(futures), total=len(futures), desc="公式商品取得"):
            section, url = futures[future]
            try:
                official_rows.extend(future.result())
            except Exception as e:
                official_errors.append({"区分": section, "URL": url, "エラー": str(e)})

    df_official = pd.DataFrame(official_rows)
    if df_official.empty:
        raise RuntimeError("公式商品を取得できませんでした。時間を置いて再実行してください。")
    df_official_errors = pd.DataFrame(official_errors)

df_official["JANコード"] = df_official["JANコード"].apply(valid_jan)
df_official = df_official.drop_duplicates(
    subset=["公式商品ページ", "JANコード"], keep="first"
).reset_index(drop=True)

all_valid_jans = sorted(df_official["JANコード"].dropna().unique().tolist())
summary_by_section = (
    df_official.groupby("公式区分", dropna=False)
    .agg(公式商品ページ数=("公式商品ページ", "nunique"), JAN数=("JANコード", "nunique"))
    .reset_index()
)

print("\n公式カタログ準備結果")
display(summary_by_section)
print(f"公式商品ページ数：{df_official['公式商品ページ'].nunique()}")
print(f"有効JAN数：{len(all_valid_jans)}")
print(f"JANなし商品ページ数：{df_official.loc[df_official['JANコード'].isna(), '公式商品ページ'].nunique()}")

if not MASTER_REUSED:
    with pd.ExcelWriter(OFFICIAL_ONLY_FILE, engine="openpyxl") as writer:
        df_official.to_excel(writer, sheet_name="公式商品マスタ", index=False)
        summary_by_section.to_excel(writer, sheet_name="区分別集計", index=False)
        df_official_errors.to_excel(writer, sheet_name="取得エラー", index=False)

    print(f"\n初回用マスタを保存しました：{OFFICIAL_ONLY_FILE}")
    print("次回からこのファイルをアップロードすれば、公式サイト巡回を省略できます。")
    if IN_COLAB:
        files.download(OFFICIAL_ONLY_FILE)


# ================================================================
# 3.5 Keepa照合範囲の指定
# ================================================================

print("\nKeepa照合範囲を選びます。")
print("AQUA=アクアリウム / CA=犬猫 / SMALL=小動物 / REPTILE=爬虫類 / ALL=全区分")
scope_text = input("対象区分（未入力はAQUA）：").strip().upper() or "AQUA"
scope_map = {
    "AQUA": ["アクアリウム"],
    "CA": ["犬猫"],
    "SMALL": ["小動物"],
    "REPTILE": ["爬虫類（エキゾテラ）"],
    "ALL": list(OFFICIAL_ROOTS.keys()),
}
selected_sections = scope_map.get(scope_text, ["アクアリウム"])

df_official_scope = df_official[df_official["公式区分"].isin(selected_sections)].copy()
filter_keyword = input(
    "商品名・カテゴリでさらに絞る場合は語句を入力（例：フィルター、未入力は絞り込みなし）："
).strip()
if filter_keyword:
    keyword_mask = (
        df_official_scope["公式商品名"].fillna("").str.contains(filter_keyword, case=False, regex=False)
        | df_official_scope["公式カテゴリ"].fillna("").str.contains(filter_keyword, case=False, regex=False)
    )
    df_official_scope = df_official_scope[keyword_mask].copy()

scope_jans = sorted(df_official_scope["JANコード"].dropna().unique().tolist())
print(f"選択範囲の有効JAN数：{len(scope_jans)}")

start_text = input("分割実行の開始位置（通常は0）：").strip()
try:
    jan_start = max(0, int(start_text or "0"))
except ValueError:
    jan_start = 0

limit_text = input(
    "今回照合する最大JAN数（推奨200、全件はALL）："
).strip().upper()
if limit_text == "ALL":
    jan_limit = None
else:
    try:
        jan_limit = max(0, int(limit_text or "200"))
    except ValueError:
        jan_limit = 200

target_jans = scope_jans[jan_start:] if jan_limit is None else scope_jans[jan_start:jan_start + jan_limit]
print(f"今回のKeepa照合対象：{len(target_jans)} JAN（開始位置 {jan_start}）")
if not target_jans:
    print("対象JANがありません。処理を終了します。")
    raise SystemExit


# ================================================================
# 4. Keepa API共通アクセス
# ================================================================

answer = input(
    "\n選択したJANをKeepaでAmazon.co.jpへ照合します。"
    f"最低約{len(target_jans)}トークンを消費します。続ける場合は YES："
).strip().upper()

if answer != "YES":
    print("Keepa照合を行わず終了します。")
    raise SystemExit

API_KEY = getpass.getpass("Keepa APIキーを入力してください：").strip()
if not API_KEY:
    raise RuntimeError("Keepa APIキーが入力されていません。")

BASE_URL = "https://api.keepa.com"
keepa_session = requests.Session()
keepa_session.headers.update({"Accept-Encoding": "gzip", "User-Agent": "GEX-Research-v4.3"})

LAST_TOKENS_LEFT = None
REFILL_RATE = None
TOTAL_TOKENS_CONSUMED = 0
keepa_errors = []


def wait_with_progress(wait_sec, reason):
    """長い待機中も1分ごとに残り時間を表示する。"""
    remaining = max(1, int(math.ceil(wait_sec)))
    print(f"\n{reason}")
    print(f"待機予定：約{math.ceil(remaining / 60)}分（{remaining}秒）")
    while remaining > 0:
        step = min(60, remaining)
        time.sleep(step)
        remaining -= step
        if remaining > 0:
            print(f"  トークン補充待ち：残り約{math.ceil(remaining / 60)}分")
    print("トークン待機が完了しました。処理を再開します。")


def ensure_estimated_tokens(required_tokens, process_name):
    """直前レスポンスの残量と補充速度から、次の処理に必要な分を待つ。"""
    global LAST_TOKENS_LEFT
    if LAST_TOKENS_LEFT is None:
        return

    target = max(1, int(required_tokens)) + TOKEN_WAIT_BUFFER
    if LAST_TOKENS_LEFT >= target:
        return

    rate = REFILL_RATE if REFILL_RATE and REFILL_RATE > 0 else DEFAULT_REFILL_RATE
    shortage = target - LAST_TOKENS_LEFT
    wait_sec = math.ceil(shortage / rate * 60) + 5
    wait_with_progress(
        wait_sec,
        f"{process_name}に必要なトークンが不足しています。"
        f"現在 {LAST_TOKENS_LEFT} / 目安 {target} / 補充速度 {rate}トークン/分",
    )
    # 待機時間から見込まれる補充分を反映し、同じ残量で再待機し続けないようにする。
    LAST_TOKENS_LEFT = target


class KeepaTokenShortage(RuntimeError):
    """Keepaのトークン補充を待てば再実行できるエラー。"""


def is_token_shortage_error(error):
    message = str(error).lower()
    markers = (
        "not enough token",
        "insufficient token",
        "token limit",
        "tokens left",
        "refill",
        "too many requests",
        "rate limit",
        "status code: 429",
        "429 client error",
        "トークン不足",
    )
    return any(marker in message for marker in markers)


def keepa_get(endpoint, params, retry=5):
    global LAST_TOKENS_LEFT, REFILL_RATE, TOTAL_TOKENS_CONSUMED
    p = params.copy()
    p["key"] = API_KEY
    error_attempts = 0

    while True:
        try:
            response = keepa_session.get(BASE_URL + endpoint, params=p, timeout=150)
            if response.status_code == 429:
                try:
                    error_data = response.json()
                except Exception:
                    error_data = {}
                LAST_TOKENS_LEFT = error_data.get("tokensLeft", LAST_TOKENS_LEFT)
                REFILL_RATE = error_data.get("refillRate", REFILL_RATE)
                raise KeepaTokenShortage(
                    f"HTTP 429: {error_data.get('error') or 'Keepa API token shortage'}"
                )

            response.raise_for_status()
            data = response.json()
            if data.get("error"):
                message = str(data["error"])
                if is_token_shortage_error(message):
                    raise KeepaTokenShortage(message)
                raise RuntimeError(message)

            LAST_TOKENS_LEFT = data.get("tokensLeft", LAST_TOKENS_LEFT)
            REFILL_RATE = data.get("refillRate", REFILL_RATE)
            consumed = data.get("tokensConsumed", 0) or 0
            TOTAL_TOKENS_CONSUMED += consumed
            print(
                f"残りトークン：{LAST_TOKENS_LEFT} / 今回消費：{consumed}"
                + (f" / 補充速度：{REFILL_RATE}/分" if REFILL_RATE is not None else "")
            )
            return data
        except KeepaTokenShortage:
            raise
        except Exception as e:
            error_attempts += 1
            if error_attempts >= retry:
                raise RuntimeError(str(e))
            wait_sec = 8 * error_attempts
            print(f"Keepa通信エラー：{e} / {wait_sec}秒後に再試行します。")
            time.sleep(wait_sec)


def keepa_get_with_wait(endpoint, params, estimated_tokens, process_name):
    """トークン不足なら補充まで待機し、同じリクエストを必ず再開する。"""
    global LAST_TOKENS_LEFT

    target = max(1, int(estimated_tokens)) + TOKEN_WAIT_BUFFER
    while True:
        ensure_estimated_tokens(estimated_tokens, process_name)
        try:
            return keepa_get(endpoint, params)
        except Exception as e:
            if not isinstance(e, KeepaTokenShortage) and not is_token_shortage_error(e):
                raise

            rate = REFILL_RATE if REFILL_RATE and REFILL_RATE > 0 else DEFAULT_REFILL_RATE
            current = LAST_TOKENS_LEFT if LAST_TOKENS_LEFT is not None else 0
            shortage = max(1, target - current)
            wait_sec = max(60, math.ceil(shortage / rate * 60) + 5)
            wait_with_progress(
                wait_sec,
                f"{process_name}中にKeepaのトークンが不足しました。"
                f"現在 {current} / 目安 {target} / 補充速度 {rate}トークン/分。"
                "補充後、同じ位置から自動再開します。",
            )
            LAST_TOKENS_LEFT = target


def product_codes(product):
    values = []
    for key in ("eanList", "upcList"):
        raw = product.get(key) or []
        if isinstance(raw, list):
            values.extend(str(x) for x in raw if x)
    return set(values)


def basic_product_fields(product):
    stats = product.get("stats") or {}
    current = stats.get("current") or []
    avg30 = stats.get("avg30") or []
    avg90 = stats.get("avg90") or []
    amazon_price = yen_from_keepa(safe_array_get(current, 0))
    new_price = yen_from_keepa(safe_array_get(current, 1))

    return {
        "ASIN": product.get("asin"),
        "Keepa商品名": product.get("title"),
        "Keepaブランド": product.get("brand"),
        "Keepaメーカー": product.get("manufacturer"),
        "Amazon現在": "あり" if amazon_price is not None else "不在",
        "Amazon現在価格": amazon_price,
        "新品現在価格": new_price,
        "現在ランキング": safe_array_get(current, 3),
        "30日平均ランキング": safe_array_get(avg30, 3),
        "90日平均ランキング": safe_array_get(avg90, 3),
        "Keepa月間販売表示": product.get("monthlySold"),
        "Amazon商品ページ": (
            "https://www.amazon.co.jp/dp/" + str(product.get("asin"))
            if product.get("asin") else None
        ),
    }


# ================================================================
# 5. 全公式JAN → Amazon ASIN照合（検索語は使わない）
# ================================================================

print("\n" + "=" * 72)
print("STEP 2：公式JANをKeepaへ照合")
print("=" * 72)

jan_to_products = defaultdict(dict)
attempted_jans = set()
basic_stopped_early = False

for start in tqdm(range(0, len(target_jans), KEEPA_CODE_BATCH), desc="JAN照合"):
    chunk = target_jans[start:start + KEEPA_CODE_BATCH]
    try:
        # JANから複数ASINが返る場合を考慮し、2トークン/件で保守的に待つ。
        data = keepa_get_with_wait("/product", {
            "domain": DOMAIN,
            "code": ",".join(chunk),
            "code-limit": 10,
            "stats": 90,
            "history": 0,
        }, len(chunk) * 2, "JAN照合")
    except Exception as e:
        keepa_errors.append({"工程": "JAN照合", "対象": ",".join(chunk), "エラー": str(e)})
        basic_stopped_early = True
        raise RuntimeError(f"JAN照合でトークン不足以外のエラーが発生しました：{e}") from e

    attempted_jans.update(chunk)

    normalized_chunk = {normalize_code(x): x for x in chunk}
    for product in data.get("products", []):
        for code in product_codes(product):
            original_jan = normalized_chunk.get(normalize_code(code))
            if original_jan:
                asin = product.get("asin")
                if asin:
                    jan_to_products[original_jan][asin] = product

matched_rows = []
for _, official in df_official_scope.iterrows():
    jan = official.get("JANコード")
    products = jan_to_products.get(jan, {}) if jan else {}
    for product in products.values():
        matched_rows.append({**official.to_dict(), **basic_product_fields(product)})

df_matched_all = pd.DataFrame(matched_rows)
df_asin_duplicates = pd.DataFrame()

if not df_matched_all.empty:
    # まず完全に同じ照合行を除外し、その後ASIN単位で必ず1行に統一する。
    df_matched_all = df_matched_all.drop_duplicates(
        subset=["公式商品ページ", "JANコード", "ASIN"], keep="first"
    ).reset_index(drop=True)
    matched_jans = set(df_matched_all["JANコード"].dropna())

    duplicate_mask = df_matched_all.duplicated(subset=["ASIN"], keep="first")
    df_asin_duplicates = df_matched_all.loc[duplicate_mask].copy()
    if not df_asin_duplicates.empty:
        first_rows = (
            df_matched_all.loc[~duplicate_mask, ["ASIN", "公式商品名", "JANコード"]]
            .drop_duplicates("ASIN")
            .set_index("ASIN")
        )
        df_asin_duplicates["重複除外理由"] = "同一ASINが既にAmazon照合済みに存在"
        df_asin_duplicates["採用された公式商品名"] = df_asin_duplicates["ASIN"].map(
            first_rows["公式商品名"]
        )
        df_asin_duplicates["採用されたJANコード"] = df_asin_duplicates["ASIN"].map(
            first_rows["JANコード"]
        )

    df_matched = (
        df_matched_all.loc[~duplicate_mask]
        .reset_index(drop=True)
    )
else:
    df_matched = df_matched_all.copy()
    matched_jans = set()

df_unmatched = df_official_scope[
    df_official_scope["JANコード"].isna()
    | (
        df_official_scope["JANコード"].isin(attempted_jans)
        & ~df_official_scope["JANコード"].isin(matched_jans)
    )
].copy()
df_unmatched["未照合理由"] = np.select(
    [
        df_unmatched["JANコード"].isna(),
        df_unmatched["JANコード"].isin(attempted_jans),
    ],
    [
        "公式ページにJANなし",
        "Keepa/Amazon.co.jpでASIN未照合",
    ],
    default="未照合",
)

df_not_scanned = df_official[
    df_official["JANコード"].notna()
    & ~df_official["JANコード"].isin(attempted_jans)
].copy()
df_not_scanned["未調査理由"] = np.where(
    df_not_scanned["公式区分"].isin(selected_sections),
    "今回の開始位置・件数範囲外、またはトークン残量により未実行",
    "今回選択していない公式区分",
)

print(f"\nAmazon照合ASIN行数：{len(df_matched)}")
print(f"同一ASINの重複除外行数：{len(df_asin_duplicates)}")
print(f"照合を試した公式JAN数：{len(attempted_jans)} / 今回対象{len(target_jans)}")
print(f"照合できた公式JAN数：{len(matched_jans)} / 試行{len(attempted_jans)}")
print(f"未照合公式商品行数：{len(df_unmatched)}")


# ================================================================
# 6. Amazonあり商品のBuy Box深掘り
# ================================================================

bb_results = []
bb_stopped_early = False

if df_matched.empty:
    print("Amazon照合結果がないためBuy Box調査はスキップします。")
    df_matched["判定"] = None
    df_matched["候補"] = None
else:
    # ASIN単位で優先順位を付ける。月販表示が高い順、次にランキングが良い順。
    asin_priority = (
        df_matched[df_matched["Amazon現在"] == "あり"]
        .drop_duplicates("ASIN")
        .copy()
    )
    asin_priority["_月販"] = pd.to_numeric(
        asin_priority["Keepa月間販売表示"], errors="coerce"
    ).fillna(-1)
    asin_priority["_順位"] = pd.to_numeric(
        asin_priority["現在ランキング"], errors="coerce"
    ).fillna(10**12)
    asin_priority = asin_priority.sort_values(
        ["_月販", "_順位"], ascending=[False, True]
    )
    all_bb_asins = asin_priority["ASIN"].dropna().astype(str).tolist()

    print("\n" + "=" * 72)
    print("STEP 3：Amazon在庫あり商品のBuy Box深掘り")
    print("=" * 72)
    print(f"深掘り可能ASIN：{len(all_bb_asins)}")
    print("Buy Box取得は1 ASINあたり基本約3トークンです。")
    print("販売数表示が多い順に処理します。")
    limit_text = input(
        "深掘り件数を入力（推奨200、全件は ALL、スキップは 0）："
    ).strip().upper()

    if limit_text == "ALL":
        bb_asins = all_bb_asins
    else:
        try:
            limit = int(limit_text or "200")
        except ValueError:
            limit = 200
        bb_asins = all_bb_asins[:max(0, limit)]

    print(f"深掘り対象：{len(bb_asins)} ASIN / 推定最大約{len(bb_asins) * 3}トークン")
    confirm_bb = "NO"
    if bb_asins:
        confirm_bb = input("この件数で実行する場合は YES：").strip().upper()

    raw_bb_products = []
    if confirm_bb == "YES":
        for start in tqdm(range(0, len(bb_asins), KEEPA_ASIN_BATCH), desc="Buy Box取得"):
            chunk = bb_asins[start:start + KEEPA_ASIN_BATCH]
            try:
                data = keepa_get_with_wait("/product", {
                    "domain": DOMAIN,
                    "asin": ",".join(chunk),
                    "buybox": 1,
                    "stats": 90,
                    "days": 90,
                }, len(chunk) * 3, "Buy Box取得")
                raw_bb_products.extend(data.get("products", []))
            except Exception as e:
                keepa_errors.append({"工程": "Buy Box", "対象": ",".join(chunk), "エラー": str(e)})
                bb_stopped_early = True
                raise RuntimeError(
                    f"Buy Box取得でトークン不足以外のエラーが発生しました：{e}"
                ) from e

    else:
        raw_bb_products = []

    # Keepa時刻
    KEEPA_OFFSET = 21564000

    def now_keepa_minutes():
        return int(time.time() // 60) - KEEPA_OFFSET

    def parse_buybox_history(history):
        if not history:
            return []
        events = []
        for i in range(0, len(history) - 1, 2):
            try:
                events.append((int(history[i]), str(history[i + 1])))
            except Exception:
                pass
        return sorted(events, key=lambda x: x[0])

    def calculate_bb_share(history, days, amazon_ids):
        events = parse_buybox_history(history)
        empty = {
            "amazon_pct": np.nan,
            "other_pct": np.nan,
            "suppressed_pct": np.nan,
            "unknown_pct": np.nan,
            "other_days": 0.0,
            "coverage_pct": 0.0,
        }
        if not events:
            return empty

        now_k = now_keepa_minutes()
        window_start = now_k - days * 24 * 60
        current_seller = None
        observation_start = None

        for ts, seller in events:
            if ts <= window_start:
                current_seller = seller
                observation_start = window_start
            else:
                break

        if current_seller is None:
            future = [e for e in events if window_start <= e[0] <= now_k]
            if not future:
                return empty
            observation_start, current_seller = future[0]

        window_events = [
            (ts, seller) for ts, seller in events
            if observation_start < ts <= now_k
        ]
        totals = {"amazon": 0, "other": 0, "suppressed": 0, "unknown": 0}

        def category(seller):
            if seller in amazon_ids:
                return "amazon"
            if seller == "-1":
                return "suppressed"
            if seller in ("-2", "", "None", "nan"):
                return "unknown"
            return "other"

        previous = observation_start
        seller = current_seller
        for ts, new_seller in window_events:
            totals[category(seller)] += max(0, ts - previous)
            previous, seller = ts, new_seller
        totals[category(seller)] += max(0, now_k - previous)

        observed = sum(totals.values())
        if observed <= 0:
            return empty
        return {
            "amazon_pct": round(totals["amazon"] / observed * 100, 1),
            "other_pct": round(totals["other"] / observed * 100, 1),
            "suppressed_pct": round(totals["suppressed"] / observed * 100, 1),
            "unknown_pct": round(totals["unknown"] / observed * 100, 1),
            "other_days": round(totals["other"] / 1440, 1),
            "coverage_pct": round(min(100, observed / (days * 1440) * 100), 1),
        }

    amazon_ids = {AMAZON_JP_SELLER_ID}
    for product in raw_bb_products:
        stats = product.get("stats") or {}
        if stats.get("buyBoxIsAmazon") is True and stats.get("buyBoxSellerId"):
            amazon_ids.add(str(stats.get("buyBoxSellerId")))

    for product in raw_bb_products:
        stats = product.get("stats") or {}
        seller = stats.get("buyBoxSellerId")
        is_amazon = stats.get("buyBoxIsAmazon")
        if is_amazon is True:
            current_type = "Amazon"
        elif seller in (None, "-1"):
            current_type = "カートなし"
        elif seller == "-2":
            current_type = "セラー不明"
        else:
            current_type = "他セラー"

        bb30 = calculate_bb_share(product.get("buyBoxSellerIdHistory"), 30, amazon_ids)
        bb90 = calculate_bb_share(product.get("buyBoxSellerIdHistory"), 90, amazon_ids)
        bb_price = stats.get("buyBoxPrice")
        bb_shipping = stats.get("buyBoxShipping")
        bb_total = None
        if bb_price is not None and bb_price >= 0:
            bb_total = yen_from_keepa(bb_price + max(0, bb_shipping or 0))

        bb_results.append({
            "ASIN": product.get("asin"),
            "現在カート": current_type,
            "現在カートSellerID": seller,
            "現在カート価格": bb_total,
            "30日_Amazonカート率%": bb30["amazon_pct"],
            "30日_他セラーカート率%": bb30["other_pct"],
            "90日_Amazonカート率%": bb90["amazon_pct"],
            "90日_他セラーカート率%": bb90["other_pct"],
            "90日_他セラーカート日数": bb90["other_days"],
            "90日_カートなし率%": bb90["suppressed_pct"],
            "90日_不明率%": bb90["unknown_pct"],
            "90日_履歴カバー率%": bb90["coverage_pct"],
        })

    df_bb = pd.DataFrame(bb_results)
    if not df_bb.empty:
        df_bb = df_bb.drop_duplicates("ASIN")
        df_matched = df_matched.merge(df_bb, on="ASIN", how="left")
    else:
        for col in [
            "現在カート", "現在カートSellerID", "現在カート価格",
            "30日_Amazonカート率%", "30日_他セラーカート率%",
            "90日_Amazonカート率%", "90日_他セラーカート率%",
            "90日_他セラーカート日数", "90日_カートなし率%",
            "90日_不明率%", "90日_履歴カバー率%",
        ]:
            df_matched[col] = np.nan

    def judge_row(row):
        if row.get("Amazon現在") == "不在":
            return "A：Amazon現在不在", "○"
        current_cart = row.get("現在カート")
        if pd.isna(current_cart):
            return "D：Buy Box未確認", ""
        if current_cart == "他セラー":
            return "B：Amazonあり・現在他セラーカート", "○"
        other90 = safe_number(row.get("90日_他セラーカート率%"))
        coverage = safe_number(row.get("90日_履歴カバー率%")) or 0
        if other90 is not None and other90 >= MIN_OTHER_BB_SHARE_90:
            return "C：Amazonあり・他セラーカート実績あり", "○"
        if coverage < MIN_BB_COVERAGE_PERCENT:
            return "F：Buy Box履歴データ不足", ""
        return "E：Amazon優勢", ""

    judged = df_matched.apply(judge_row, axis=1, result_type="expand")
    df_matched["判定"] = judged[0]
    df_matched["候補"] = judged[1]


# ================================================================
# 7. 結果整理とExcel出力
# ================================================================

if "判定" not in df_matched.columns:
    df_matched["判定"] = None
    df_matched["候補"] = None

judgment_order = {
    "A：Amazon現在不在": 1,
    "B：Amazonあり・現在他セラーカート": 2,
    "C：Amazonあり・他セラーカート実績あり": 3,
    "F：Buy Box履歴データ不足": 4,
    "D：Buy Box未確認": 5,
    "E：Amazon優勢": 6,
}

if not df_matched.empty:
    df_matched["_判定順"] = df_matched["判定"].map(judgment_order).fillna(99)
    df_matched["_月販"] = pd.to_numeric(
        df_matched["Keepa月間販売表示"], errors="coerce"
    ).fillna(-1)
    df_matched = (
        df_matched.sort_values(["_判定順", "_月販"], ascending=[True, False])
        .drop(columns=["_判定順", "_月販"])
        .reset_index(drop=True)
    )

df_candidates = df_matched[df_matched["候補"] == "○"].copy()
df_bb_unchecked = df_matched[df_matched["判定"] == "D：Buy Box未確認"].copy()
df_keepa_errors = pd.DataFrame(keepa_errors)

# 出力直前にもASIN一意性を検証し、将来の修正で重複が戻った場合は明示的に止める。
if not df_matched.empty and df_matched["ASIN"].duplicated().any():
    raise RuntimeError("Amazon照合済みにASIN重複が残っています。")
if not df_candidates.empty and df_candidates["ASIN"].duplicated().any():
    raise RuntimeError("候補_全てにASIN重複が残っています。")

counts = {
    "公式商品ページ数": int(df_official["公式商品ページ"].nunique()),
    "公式有効JAN数（全区分）": int(len(all_valid_jans)),
    "今回選択した区分": "、".join(selected_sections),
    "今回選択範囲のJAN数": int(len(scope_jans)),
    "今回の照合対象JAN数": int(len(target_jans)),
    "実際に照合を試したJAN数": int(len(attempted_jans)),
    "Amazon照合ASIN行数": int(len(df_matched)),
    "同一ASINの重複除外行数": int(len(df_asin_duplicates)),
    "照合済み公式JAN数": int(len(matched_jans)),
    "Keepa未照合商品行数": int(len(df_unmatched)),
    "候補合計": int(len(df_candidates)),
    "A Amazon現在不在": int((df_matched["判定"] == "A：Amazon現在不在").sum()),
    "B 現在他セラーカート": int((df_matched["判定"] == "B：Amazonあり・現在他セラーカート").sum()),
    "C 他セラーカート実績": int((df_matched["判定"] == "C：Amazonあり・他セラーカート実績あり").sum()),
    "D Buy Box未確認": int((df_matched["判定"] == "D：Buy Box未確認").sum()),
    "E Amazon優勢": int((df_matched["判定"] == "E：Amazon優勢").sum()),
    "F 履歴データ不足": int((df_matched["判定"] == "F：Buy Box履歴データ不足").sum()),
    "Keepa消費トークン": int(TOTAL_TOKENS_CONSUMED),
    "Keepa残りトークン": LAST_TOKENS_LEFT,
}

summary_rows = [{"項目": k, "値": v} for k, v in counts.items()]
summary_rows.extend([
    {"項目": "母集団", "値": "GEX公式商品サイト4区分（商品名検索は不使用）"},
    {"項目": "公式マスタ", "値": "既存Excelを再利用" if MASTER_REUSED else "公式サイトから新規作成"},
    {"項目": "Amazon照合キー", "値": "公式JANコード"},
    {"項目": "分割実行", "値": f"開始位置={jan_start} / 最大件数={'ALL' if jan_limit is None else jan_limit}"},
    {"項目": "追加絞り込み", "値": filter_keyword or "なし"},
    {"項目": "注意1", "値": "Keepa未照合はAmazonに商品がないと断定する意味ではありません。JAN未登録・別JAN・セット品等を含みます。"},
    {"項目": "注意2", "値": "エキゾテラはGEX公式の爬虫類用品ブランドとして対象に含めています。"},
    {"項目": "注意3", "値": "Amazon照合済み・候補_全てはASINごとに1行です。同一ASINの除外行はASIN重複除外タブで確認できます。"},
    {"項目": "実行日時UTC", "値": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")},
])
df_summary = pd.DataFrame(summary_rows)

with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    df_summary.to_excel(writer, sheet_name="サマリー", index=False)
    df_candidates.to_excel(writer, sheet_name="候補_全て", index=False)
    df_matched.to_excel(writer, sheet_name="Amazon照合済み", index=False)
    df_official.to_excel(writer, sheet_name="公式商品マスタ", index=False)
    df_unmatched.to_excel(writer, sheet_name="Keepa未照合", index=False)
    df_not_scanned.to_excel(writer, sheet_name="今回未調査", index=False)
    df_bb_unchecked.to_excel(writer, sheet_name="BB未確認", index=False)
    df_asin_duplicates.to_excel(writer, sheet_name="ASIN重複除外", index=False)
    df_official_errors.to_excel(writer, sheet_name="公式取得エラー", index=False)
    df_keepa_errors.to_excel(writer, sheet_name="Keepa取得エラー", index=False)


# ================================================================
# 8. Excel装飾
# ================================================================

wb = load_workbook(OUTPUT_FILE)
header_fill = PatternFill("solid", fgColor="1F4E78")
candidate_fill = PatternFill("solid", fgColor="E2F0D9")
warning_fill = PatternFill("solid", fgColor="FFF2CC")
link_fill = PatternFill("solid", fgColor="DDEBF7")

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = {cell.value: cell.column for cell in ws[1]}
    candidate_col = headers.get("候補")
    judgment_col = headers.get("判定")

    for row in range(2, ws.max_row + 1):
        if candidate_col and ws.cell(row, candidate_col).value == "○":
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col).fill = candidate_fill
        if judgment_col:
            value = str(ws.cell(row, judgment_col).value or "")
            if "未確認" in value or "データ不足" in value:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row, col).fill = warning_fill

        for url_header in ("公式商品ページ", "Amazon商品ページ"):
            col = headers.get(url_header)
            if col:
                cell = ws.cell(row, col)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
                    cell.fill = link_fill

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, min(ws.max_row, 300) + 1):
            value = ws.cell(row, col_idx).value
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)

wb.save(OUTPUT_FILE)


# ================================================================
# 9. 完了表示・ダウンロード
# ================================================================

print("\n" + "=" * 72)
print("GEX公式カタログ起点 Keepa調査 v4.3 完了")
print("=" * 72)
for key, value in counts.items():
    print(f"{key}：{value}")

display_cols = [
    "判定", "公式区分", "公式商品名", "JANコード", "ASIN",
    "Amazon現在", "Amazon現在価格", "現在カート",
    "30日_他セラーカート率%", "90日_他セラーカート率%",
    "Keepa月間販売表示", "現在ランキング",
]
display_cols = [c for c in display_cols if c in df_candidates.columns]
if len(df_candidates):
    print("\n===== 候補商品（先頭100件） =====")
    display(df_candidates[display_cols].head(100))

print(f"\n出力ファイル：{OUTPUT_FILE}")
if IN_COLAB:
    files.download(OUTPUT_FILE)
